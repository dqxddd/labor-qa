"""资料解析：把上传的文件读成纯文本。

9.17 先只支持 md / txt。
9.22 补上 PDF 和 Word：
  - PDF 用 pypdf（纯 Python，不用编译，装上就能用）
  - Word 用 python-docx（只认 .docx，老版 .doc 要先另存为 .docx）
这两个库干的都是「把文字抠出来」，遇到扫描件（整页是图片、没有文字层）
抠不出东西，会明确报错，不会静默返回空内容。
"""

import os
import re


class ParseError(Exception):
    pass


SUPPORTED_TYPES = {"md", "txt", "pdf", "docx"}

# 中文名，报错信息和前端提示都能用
TYPE_LABELS = {"md": "Markdown", "txt": "纯文本", "pdf": "PDF", "docx": "Word"}

# 老版 .doc 是二进制格式，python-docx 读不了，单独给一句人话的提示
_HINTS = {"doc": "老版 .doc 解析不了，用 Word 打开「另存为 .docx」再传"}

# 有些 PDF 抽字的时候会在每个汉字之间塞一个空格（「劳 动 合 同」），
# 检索是按连续两个字比对的，被切断就搜不到了，得清掉。
# 但正常文档里本来就有零星空格（「第一条 为了……」），清掉反而不好看，
# 所以先数一下密度，特别密的才当毛病处理。只动空格和制表符，不动换行。
_CJK = r"\u4e00-\u9fff\u3000-\u303f\uff01-\uff60"
_CJK_SPACE = re.compile(r"(?<=[%s])[ \t\u00a0]+(?=[%s])" % (_CJK, _CJK))
_CJK_CHAR = re.compile(r"[%s]" % _CJK)
_CJK_SPACE_RATIO = 0.15      # 汉字间空格占比超过这个数，就认定是抽字的毛病


def load_text(file_path, file_type):
    """返回文件里的纯文本内容，失败时抛 ParseError"""
    file_type = (file_type or "").lower().lstrip(".")

    if file_type not in SUPPORTED_TYPES:
        hint = _HINTS.get(file_type)
        if hint:
            raise ParseError(hint)
        raise ParseError("不支持 %s 格式，目前能解析：%s"
                         % (file_type or "未知", "、".join(sorted(SUPPORTED_TYPES))))

    if not os.path.exists(file_path):
        raise ParseError("文件不存在：%s" % file_path)

    if file_type == "pdf":
        text = _read_pdf(file_path)
    elif file_type == "docx":
        text = _read_docx(file_path)
    else:
        text = _read_plain(file_path)

    text = _tidy(text)
    if not text:
        raise ParseError("文件里没有文字内容")

    return text


def _read_plain(file_path):
    """md / txt：先按 utf-8 读，读不了再试 gbk（国内不少 txt 是 gbk）"""
    try:
        with open(file_path, "rb") as f:
            raw = f.read()
    except OSError as e:
        raise ParseError("读文件失败：%s" % e)

    for encoding in ("utf-8", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue

    # 两种都不对，退回「读不出来的字节直接丢掉」，至少还能用
    return raw.decode("utf-8", errors="ignore")


def _read_pdf(file_path):
    """PDF：用 pypdf 逐页抠文字"""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ParseError("没装 pypdf，解析不了 PDF。在项目目录执行：pip install pypdf")

    try:
        reader = PdfReader(file_path)
        encrypted = reader.is_encrypted
        if encrypted:
            # 很多 PDF 只是设了「权限口令」，打开口令为空，试一下
            try:
                opened = bool(reader.decrypt(""))
            except Exception:
                opened = False
            if not opened:
                raise ParseError("这份 PDF 有密码，解不开。先去密码再传。")
        pages = list(reader.pages)
    except ParseError:
        raise
    except Exception as e:
        raise ParseError("PDF 打不开，文件可能损坏：%s" % e)

    texts = []
    for page in pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")        # 某一页坏掉不影响整份文件

    text = "\n".join(texts)

    if not text.strip():
        raise ParseError("这份 PDF 一页文字都没抠出来，多半是扫描件（整页是图片）。"
                         "请换成带文字层的 PDF，或者先 OCR 一遍再传。")

    return text


def _read_docx(file_path):
    """Word(.docx)：段落 + 表格里的文字都抠出来"""
    try:
        import docx
    except ImportError:
        raise ParseError("没装 python-docx，解析不了 Word。在项目目录执行：pip install python-docx")

    try:
        document = docx.Document(file_path)
        blocks = [p.text for p in document.paragraphs]
        tables = document.tables
    except Exception as e:
        raise ParseError("Word 打不开，确认是 .docx 格式：%s" % e)

    # 表格里的字也要 —— 劳动合同、工资标准这类内容经常是放在表格里的
    for table in tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                value = cell.text.strip()
                # 合并单元格会让相邻格子的文字重复，去个重
                if value and (not cells or cells[-1] != value):
                    cells.append(value)
            if cells:
                blocks.append(" | ".join(cells))

    return "\n".join(blocks)


def _tidy(text):
    """统一清理：修掉 PDF 抽字留下的汉字间空格，把连续空行压成一个"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _drop_cjk_spaces(text)

    # 连续空行压成一个，不然切分出来的片段全是空行
    lines = [line.rstrip() for line in text.splitlines()]
    cleaned = []
    for line in lines:
        if line == "" and cleaned and cleaned[-1] == "":
            continue
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def _drop_cjk_spaces(text):
    """汉字之间的空格太多了就整篇清掉，只有零星几个就留着"""
    cjk_count = len(_CJK_CHAR.findall(text))
    if cjk_count < 40:
        return text

    gaps = len(_CJK_SPACE.findall(text))
    if gaps / float(cjk_count) < _CJK_SPACE_RATIO:
        return text          # 零星空格，是正常写法，别动

    return _CJK_SPACE.sub("", text)
